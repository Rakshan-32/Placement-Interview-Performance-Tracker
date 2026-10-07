"""
User access upload tests — updated for the invitation-based lifecycle.
New users are created as INVITED (not immediately ACTIVE).
"""
import openpyxl
import io
import db
from fastapi.testclient import TestClient
from app import app

client = TestClient(app)


def _coord_headers():
    coord = db.get_user_by_gmail("coordinator@gmail.com")
    return {"X-User-Id": coord["uuid"]}


def test_bulk_user_access_upload():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "User Access"
    ws.append(["gmail", "role"])
    ws.append(["test_student1@gmail.com", "Student"])
    ws.append(["test_recruiter1@gmail.com", "Recruiter"])
    ws.append(["test_coord1@gmail.com", "Coordinator"])

    excel_file = io.BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)

    response = client.post(
        "/api/users/upload-access",
        data={"default_role": "Student"},
        files={"file": ("test_users.xlsx", excel_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_coord_headers()
    )

    assert response.status_code == 200, f"Response failed: {response.text}"
    json_data = response.json()
    assert json_data["success"] is True
    assert json_data["total_processed"] == 3
    assert json_data["invited_count"] == 3

    student = db.get_user_by_gmail("test_student1@gmail.com")
    assert student is not None
    assert student["role"] == "Student"
    assert student["access_status"] == "INVITED"
    assert student["is_active"] == 0

    recruiter = db.get_user_by_gmail("test_recruiter1@gmail.com")
    assert recruiter is not None
    assert recruiter["role"] == "Recruiter"

    coord = db.get_user_by_gmail("test_coord1@gmail.com")
    assert coord is not None
    assert coord["role"] == "Coordinator"


def test_bulk_upload_role_update_for_active_user():
    """After activating a user, bulk upload can update their role."""
    coord = db.get_user_by_gmail("coordinator@gmail.com")
    result = db.grant_single_user_access(
        "bulk_role_update@gmail.com", "Student",
        actor_uuid=coord["uuid"], actor_gmail=coord["gmail"]
    )
    token = result.get("activation_token")
    if token:
        db.activate_user(token, "TestPass123")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["gmail", "role"])
    ws.append(["bulk_role_update@gmail.com", "Recruiter"])

    excel_file = io.BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)

    res = client.post(
        "/api/users/upload-access",
        data={"default_role": "Student"},
        files={"file": ("test_users2.xlsx", excel_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_coord_headers()
    )
    assert res.status_code == 200
    json_data = res.json()
    assert json_data["updated_count"] >= 1

    updated_user = db.get_user_by_gmail("bulk_role_update@gmail.com")
    assert updated_user["role"] == "Recruiter"


def test_single_user_access_invitation():
    """Single user grant creates an INVITED user (no password in request)."""
    res = client.post(
        "/api/users/grant-single-access",
        json={"gmail": "single_student@gmail.com", "role": "Student"},
        headers=_coord_headers()
    )
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert "password" not in str(data["user"]), "Password must not be exposed in API response"

    user_in_db = db.get_user_by_gmail("single_student@gmail.com")
    assert user_in_db is not None
    assert user_in_db["access_status"] == "INVITED"
    assert user_in_db["is_active"] == 0


def test_bulk_upload_skips_invited_users():
    """Bulk upload should skip already-invited users."""
    db.grant_single_user_access("bulk_invited_skip@gmail.com", "Student")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["gmail", "role"])
    ws.append(["bulk_invited_skip@gmail.com", "Student"])

    excel_file = io.BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)

    res = client.post(
        "/api/users/upload-access",
        data={"default_role": "Student"},
        files={"file": ("users_skip.xlsx", excel_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_coord_headers()
    )
    assert res.status_code == 200
    data = res.json()
    assert data["skipped_invited"] >= 1


if __name__ == "__main__":
    test_bulk_user_access_upload()
    test_bulk_upload_role_update_for_active_user()
    test_single_user_access_invitation()
    test_bulk_upload_skips_invited_users()
