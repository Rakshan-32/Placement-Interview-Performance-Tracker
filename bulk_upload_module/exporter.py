import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any, Tuple
from fastapi.responses import Response

def style_worksheet(ws):
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = data_font
            cell.border = border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)


def build_excel_response(headers: List[str], data_rows: List[List[Any]], sheet_title: str, filename: str) -> Response:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)

    for row in data_rows:
        ws.append(row)

    style_worksheet(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


def build_csv_response(headers: List[str], data_rows: List[List[Any]], filename: str) -> Response:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(data_rows)

    return Response(
        content=buffer.getvalue().encode("utf-8-sig"),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


# ==============================================================
# 1. EXPORT COMPANY DRIVES
# ==============================================================

def export_company_drives_data(drives: List[Dict[str, Any]], format_type: str = "xlsx") -> Response:
    headers = [
        "Company Name", "Job Role", "CTC LPA", "Company Type", 
        "Required CGPA", "Allowed Branches", "Total Rounds", 
        "Location", "Drive Date", "Status", "Created At"
    ]
    data_rows = []
    for d in drives:
        data_rows.append([
            d.get("company_name", ""),
            d.get("job_role", ""),
            d.get("ctc_lpa", 0.0),
            d.get("company_type", "PRODUCT"),
            d.get("required_cgpa", 0.0),
            d.get("allowed_branches", "All"),
            d.get("total_rounds", 4),
            d.get("location", "On Campus"),
            d.get("drive_date", ""),
            d.get("status", "Active"),
            d.get("created_at", "")
        ])

    if format_type.lower() == "csv":
        return build_csv_response(headers, data_rows, "company_drives_export.csv")
    return build_excel_response(headers, data_rows, "Company Drives", "company_drives_export.xlsx")


# ==============================================================
# 2. EXPORT DRIVE RESULTS / STUDENT DRIVE CANDIDATES
# ==============================================================

def export_drive_results_data(results: List[Dict[str, Any]], drive_info: Dict[str, Any] = None, format_type: str = "xlsx") -> Response:
    company_name = drive_info.get("company_name", "Drive") if drive_info else "Drive"
    role = drive_info.get("job_role", "") if drive_info else ""
    safe_name = f"{company_name.lower().replace(' ', '_')}_results"

    headers = [
        "Student Gmail", "Current Round", "Score / Marks", "Selection Verdict / Status", "Last Updated"
    ]
    data_rows = []
    for r in results:
        data_rows.append([
            r.get("gmail", ""),
            r.get("round", 1),
            r.get("score") if r.get("score") is not None else "N/A",
            r.get("result", ""),
            r.get("updated_at", "")
        ])

    if format_type.lower() == "csv":
        return build_csv_response(headers, data_rows, f"{safe_name}.csv")
    return build_excel_response(headers, data_rows, f"{company_name[:20]} Results", f"{safe_name}.xlsx")


# ==============================================================
# 3. EXPORT STUDENT ACADEMIC ROSTER
# ==============================================================

def export_student_roster_data(students: List[Dict[str, Any]], format_type: str = "xlsx") -> Response:
    headers = [
        "Register Number", "Full Name", "Student Email", "Department", 
        "CGPA", "10th Percentage", "12th Percentage", "Technical Skills", "Registered Date"
    ]
    data_rows = []
    for s in students:
        data_rows.append([
            s.get("register_number", ""),
            s.get("name", ""),
            s.get("email", ""),
            s.get("department", ""),
            s.get("cgpa", 0.0),
            s.get("tenth_percentage", "N/A"),
            s.get("twelfth_percentage", "N/A"),
            s.get("skills", ""),
            s.get("created_at", "")
        ])

    if format_type.lower() == "csv":
        return build_csv_response(headers, data_rows, "student_roster_export.csv")
    return build_excel_response(headers, data_rows, "Student Roster", "student_roster_export.xlsx")


# ==============================================================
# 4. EXPORT USER ACCESS ACCOUNTS
# ==============================================================

def export_user_access_data(users: List[Dict[str, Any]], format_type: str = "xlsx") -> Response:
    headers = ["User Email", "Role", "Is Active", "Account Created At"]
    data_rows = []
    for u in users:
        data_rows.append([
            u.get("gmail", ""),
            u.get("role", "Student"),
            "Active" if u.get("is_active") else "Inactive",
            u.get("created_at", "")
        ])

    if format_type.lower() == "csv":
        return build_csv_response(headers, data_rows, "user_accounts_export.csv")
    return build_excel_response(headers, data_rows, "User Accounts", "user_accounts_export.xlsx")
