"""
Comprehensive tests for the Complete Access Lifecycle (Unit 1).

Covers: invitation, activation, token security, password hashing,
legacy migration, forgot/reset password, revoke, reactivate,
resend invitation, bulk invitation, authorization, and audit history.
"""
import pytest
import time
import os
import openpyxl
import io
from unittest.mock import patch, MagicMock
import smtplib
import db
import email_service
from fastapi.testclient import TestClient
from app import app

client = TestClient(app)


# ==============================================================
# HELPERS
# ==============================================================

def _coord():
    user = db.get_user_by_gmail("coordinator@gmail.com")
    assert user is not None
    return user

def _coord_headers():
    return {"X-User-Id": _coord()["uuid"]}

def _create_active_user(gmail: str, role: str = "Student"):
    """Helper: invite user via DB, then activate with a password. Returns (user, password)."""
    coord = _coord()
    result = db.grant_single_user_access(
        gmail, role,
        actor_uuid=coord["uuid"], actor_gmail=coord["gmail"]
    )
    token = result.get("activation_token")
    if token:
        password = "TestPass123"
        success, _ = db.activate_user(token, password)
        assert success, f"Activation failed for {gmail}"
        user = db.get_user_by_gmail(gmail)
        return user, password
    user = db.get_user_by_gmail(gmail)
    return user, None


# ==============================================================
# 35A: INVITATION FLOW
# ==============================================================

class TestInvitationFlow:
    def test_grant_new_user_creates_invited(self):
        res = client.post("/api/users/grant-single-access", json={
            "gmail": "lifecycle_invite1@gmail.com", "role": "Student"
        }, headers=_coord_headers())
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True
        assert "password" not in str(data["user"])

        user = db.get_user_by_gmail("lifecycle_invite1@gmail.com")
        assert user is not None
        assert user["access_status"] == "INVITED"
        assert user["is_active"] == 0

    def test_invitation_creates_history(self):
        history = db.get_access_history(user_gmail="lifecycle_invite1@gmail.com")
        invited = [h for h in history if h["action"] == "INVITED"]
        assert len(invited) >= 1
        assert invited[0]["new_status"] == "INVITED"

    @patch.dict(os.environ, {"SMTP_HOST": "", "SMTP_USERNAME": "", "SMTP_PASSWORD": "", "SMTP_FROM": ""}, clear=False)
    def test_grant_returns_dev_activation_url_when_no_smtp(self):
        res = client.post("/api/users/grant-single-access", json={
            "gmail": "lifecycle_invite_url@gmail.com", "role": "Student"
        }, headers=_coord_headers())
        data = res.json()
        assert "dev_activation_url" in data
        assert "#/activate/" in data["dev_activation_url"]

    def test_grant_already_invited_returns_409(self):
        res = client.post("/api/users/grant-single-access", json={
            "gmail": "lifecycle_invite1@gmail.com", "role": "Student"
        }, headers=_coord_headers())
        assert res.status_code == 409
        assert "already invited" in res.json()["message"].lower() or "invited" in res.json()["message"].lower()

    def test_grant_invalid_email_rejected(self):
        res = client.post("/api/users/grant-single-access", json={
            "gmail": "not-an-email", "role": "Student"
        }, headers=_coord_headers())
        assert res.status_code == 400


# ==============================================================
# 35B: INVITED USER LOGIN BLOCKING
# ==============================================================

class TestInvitedLoginBlocking:
    def test_invited_user_cannot_login(self):
        db.grant_single_user_access("blocked_invite@gmail.com", "Student")
        res = client.post("/api/login", json={
            "gmail": "blocked_invite@gmail.com", "password": "anything"
        })
        assert res.status_code == 403
        data = res.json()
        assert data["success"] is False
        assert "activation" in data["message"].lower()

    def test_invited_coordinator_blocked_from_managed(self):
        result = db.grant_single_user_access("invited_coord@gmail.com", "Coordinator")
        user = db.get_user_by_gmail("invited_coord@gmail.com")
        res = client.get("/api/users/managed", headers={"X-User-Id": user["uuid"]})
        assert res.status_code == 403

    def test_invited_user_blocked_from_interventions(self):
        result = db.grant_single_user_access("invited_mentor@gmail.com", "Mentor")
        user = db.get_user_by_gmail("invited_mentor@gmail.com")
        headers = {
            "X-User-Id": user["uuid"],
            "X-User-Role": user["role"],
            "X-Department": "CSE"
        }
        res = client.get("/api/interventions", headers=headers)
        assert res.status_code == 403
        assert "activation" in res.json()["detail"].lower()


# ==============================================================
# 35C: ACTIVATION FLOW
# ==============================================================

class TestActivationFlow:
    def test_validate_activation_token(self):
        result = db.grant_single_user_access("activate_test@gmail.com", "Student")
        token = result["activation_token"]
        res = client.get(f"/api/auth/validate-token?token={token}&purpose=ACTIVATION")
        assert res.status_code == 200
        data = res.json()
        assert data["valid"] is True
        assert data["gmail"] == "activate_test@gmail.com"
        assert data["role"] == "Student"

    def test_activate_account_success(self):
        result = db.grant_single_user_access("activate_full@gmail.com", "Student")
        token = result["activation_token"]
        res = client.post("/api/auth/activate", json={
            "token": token, "password": "SecurePass1"
        })
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True

        user = db.get_user_by_gmail("activate_full@gmail.com")
        assert user["access_status"] == "ACTIVE"
        assert user["is_active"] == 1

    def test_activated_user_can_login(self):
        res = client.post("/api/login", json={
            "gmail": "activate_full@gmail.com", "password": "SecurePass1"
        })
        assert res.status_code == 200
        assert res.json()["success"] is True

    def test_activation_creates_history(self):
        history = db.get_access_history(user_gmail="activate_full@gmail.com")
        activated = [h for h in history if h["action"] == "ACTIVATED"]
        assert len(activated) >= 1
        assert activated[0]["old_status"] == "INVITED"
        assert activated[0]["new_status"] == "ACTIVE"

    def test_short_password_rejected(self):
        result = db.grant_single_user_access("activate_short@gmail.com", "Student")
        token = result["activation_token"]
        res = client.post("/api/auth/activate", json={
            "token": token, "password": "ab"
        })
        assert res.status_code == 400
        assert "6 characters" in res.json()["message"]

    def test_used_token_rejected(self):
        result = db.grant_single_user_access("activate_used@gmail.com", "Student")
        token = result["activation_token"]
        client.post("/api/auth/activate", json={"token": token, "password": "ValidPass1"})
        res = client.post("/api/auth/activate", json={"token": token, "password": "AnotherPass"})
        assert res.status_code == 400
        assert "used" in res.json()["message"].lower()

    def test_invalid_token_rejected(self):
        res = client.post("/api/auth/activate", json={
            "token": "completely-bogus-token", "password": "ValidPass1"
        })
        assert res.status_code == 400


# ==============================================================
# 35D: TOKEN SECURITY
# ==============================================================

class TestTokenSecurity:
    def test_token_stored_as_hash(self):
        result = db.grant_single_user_access("token_hash_test@gmail.com", "Student")
        raw_token = result["activation_token"]
        conn = db.get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT token_hash FROM auth_tokens WHERE user_uuid = ?", (result["uuid"],))
        row = cursor.fetchone()
        conn.close()
        assert row is not None
        assert row["token_hash"] != raw_token
        assert len(row["token_hash"]) == 64  # SHA-256 hex

    def test_previous_tokens_invalidated(self):
        result = db.grant_single_user_access("token_invalidate@gmail.com", "Student")
        first_token = result["activation_token"]
        second_token = db.create_auth_token(result["uuid"], "ACTIVATION")
        validation1 = db.validate_auth_token(first_token, "ACTIVATION")
        assert not validation1["valid"]
        validation2 = db.validate_auth_token(second_token, "ACTIVATION")
        assert validation2["valid"]

    def test_expired_token_rejected(self):
        result = db.grant_single_user_access("token_expire@gmail.com", "Student")
        raw_token = db.create_auth_token(result["uuid"], "ACTIVATION", expiry_seconds=0)
        time.sleep(0.1)
        validation = db.validate_auth_token(raw_token, "ACTIVATION")
        assert not validation["valid"]
        assert "expired" in validation["error"].lower()


# ==============================================================
# 35E: PASSWORD HASHING
# ==============================================================

class TestPasswordHashing:
    def test_activated_password_is_bcrypt_hashed(self):
        result = db.grant_single_user_access("hash_test@gmail.com", "Student")
        db.activate_user(result["activation_token"], "MySecurePass")
        user = db.get_user_by_gmail("hash_test@gmail.com")
        assert db.is_hashed(user["password"])
        assert user["password"].startswith("$2b$")

    def test_verify_password_works_with_hash(self):
        hashed = db.hash_password("testpassword")
        assert db.verify_password("testpassword", hashed)
        assert not db.verify_password("wrongpassword", hashed)


# ==============================================================
# 35F: LEGACY ACCOUNT COMPATIBILITY
# ==============================================================

class TestLegacyAccounts:
    def test_seed_user_plaintext_login(self):
        res = client.post("/api/login", json={
            "gmail": "student@gmail.com", "password": "student123"
        })
        assert res.status_code == 200
        assert res.json()["success"] is True

    def test_legacy_password_migrated_to_bcrypt(self):
        client.post("/api/login", json={
            "gmail": "mentor@gmail.com", "password": "mentor123"
        })
        user = db.get_user_by_gmail("mentor@gmail.com")
        assert db.is_hashed(user["password"])

    def test_migrated_password_still_works(self):
        res = client.post("/api/login", json={
            "gmail": "mentor@gmail.com", "password": "mentor123"
        })
        assert res.status_code == 200


# ==============================================================
# 35G: FORGOT PASSWORD
# ==============================================================

class TestForgotPassword:
    def test_forgot_password_active_user(self):
        _create_active_user("forgot_active@gmail.com", "Student")
        res = client.post("/api/auth/forgot-password", json={
            "gmail": "forgot_active@gmail.com"
        })
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True
        assert "email_sent" in data or "dev_reset_url" in data

    def test_forgot_password_nonexistent_user_safe_response(self):
        res = client.post("/api/auth/forgot-password", json={
            "gmail": "nonexistent_forgot@gmail.com"
        })
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True
        assert "dev_reset_url" not in data

    def test_forgot_password_invited_user_no_token(self):
        db.grant_single_user_access("forgot_invited@gmail.com", "Student")
        res = client.post("/api/auth/forgot-password", json={
            "gmail": "forgot_invited@gmail.com"
        })
        data = res.json()
        assert "dev_reset_url" not in data

    def test_forgot_password_revoked_user_no_token(self):
        user, pwd = _create_active_user("forgot_revoked@gmail.com", "Student")
        db.revoke_user_access("forgot_revoked@gmail.com")
        res = client.post("/api/auth/forgot-password", json={
            "gmail": "forgot_revoked@gmail.com"
        })
        data = res.json()
        assert "dev_reset_url" not in data


# ==============================================================
# 35H: RESET PASSWORD
# ==============================================================

class TestResetPassword:
    def test_reset_password_flow(self):
        _create_active_user("reset_flow@gmail.com", "Student")
        raw_token, user = db.request_password_reset("reset_flow@gmail.com")
        assert raw_token is not None

        res = client.post("/api/auth/reset-password", json={
            "token": raw_token, "password": "NewResetPass1"
        })
        assert res.status_code == 200
        assert res.json()["success"] is True

        login_res = client.post("/api/login", json={
            "gmail": "reset_flow@gmail.com", "password": "NewResetPass1"
        })
        assert login_res.status_code == 200

    def test_old_password_fails_after_reset(self):
        login_res = client.post("/api/login", json={
            "gmail": "reset_flow@gmail.com", "password": "TestPass123"
        })
        assert login_res.status_code == 401

    def test_reset_creates_history(self):
        history = db.get_access_history(user_gmail="reset_flow@gmail.com")
        resets = [h for h in history if h["action"] == "PASSWORD_RESET"]
        assert len(resets) >= 1

    def test_reset_with_invalid_token(self):
        res = client.post("/api/auth/reset-password", json={
            "token": "bogus-token", "password": "ValidPass1"
        })
        assert res.status_code == 400

    def test_reset_with_short_password(self):
        _create_active_user("reset_short@gmail.com", "Student")
        raw_token, _ = db.request_password_reset("reset_short@gmail.com")
        res = client.post("/api/auth/reset-password", json={
            "token": raw_token, "password": "ab"
        })
        assert res.status_code == 400


# ==============================================================
# 35I: REVOKE ACCESS
# ==============================================================

class TestRevokeAccess:
    def test_revoke_active_user(self):
        _create_active_user("lc_revoke@gmail.com", "Student")
        res = client.patch("/api/users/revoke", json={
            "gmail": "lc_revoke@gmail.com"
        }, headers=_coord_headers())
        assert res.status_code == 200
        assert res.json()["user"]["is_active"] == 0

        user = db.get_user_by_gmail("lc_revoke@gmail.com")
        assert user["access_status"] == "REVOKED"
        assert user["is_active"] == 0

    def test_revoked_user_cannot_login(self):
        res = client.post("/api/login", json={
            "gmail": "lc_revoke@gmail.com", "password": "TestPass123"
        })
        assert res.status_code == 403
        assert "revoked" in res.json()["message"].lower()

    def test_revoke_creates_history(self):
        history = db.get_access_history(user_gmail="lc_revoke@gmail.com")
        revoked = [h for h in history if h["action"] == "REVOKED"]
        assert len(revoked) >= 1
        assert revoked[0]["new_status"] == "REVOKED"

    def test_revoke_nonexistent_user_404(self):
        res = client.patch("/api/users/revoke", json={
            "gmail": "nonexistent_lc@gmail.com"
        }, headers=_coord_headers())
        assert res.status_code == 404

    def test_revoke_already_revoked_safe(self):
        res = client.patch("/api/users/revoke", json={
            "gmail": "lc_revoke@gmail.com"
        }, headers=_coord_headers())
        assert res.status_code == 200
        assert "already" in res.json()["message"].lower()

    def test_cannot_revoke_self(self):
        coord = _coord()
        res = client.patch("/api/users/revoke", json={
            "gmail": coord["gmail"]
        }, headers=_coord_headers())
        assert res.status_code == 400

    def test_grant_revoked_user_returns_409(self):
        res = client.post("/api/users/grant-single-access", json={
            "gmail": "lc_revoke@gmail.com", "role": "Student"
        }, headers=_coord_headers())
        assert res.status_code == 409

    def test_student_cannot_revoke(self):
        student = db.get_user_by_gmail("student@gmail.com")
        res = client.patch("/api/users/revoke", json={
            "gmail": "lc_revoke@gmail.com"
        }, headers={"X-User-Id": student["uuid"]})
        assert res.status_code == 403


# ==============================================================
# 35J: REACTIVATE ACCESS
# ==============================================================

class TestReactivateAccess:
    def test_reactivate_revoked_user(self):
        res = client.patch("/api/users/reactivate", json={
            "gmail": "lc_revoke@gmail.com"
        }, headers=_coord_headers())
        assert res.status_code == 200
        data = res.json()
        assert data["user"]["is_active"] == 1
        assert data["user"]["access_status"] == "ACTIVE"

    def test_reactivated_user_can_login(self):
        res = client.post("/api/login", json={
            "gmail": "lc_revoke@gmail.com", "password": "TestPass123"
        })
        assert res.status_code == 200

    def test_reactivate_creates_history(self):
        history = db.get_access_history(user_gmail="lc_revoke@gmail.com")
        reactivated = [h for h in history if h["action"] == "REACTIVATED"]
        assert len(reactivated) >= 1

    def test_reactivate_already_active_safe(self):
        res = client.patch("/api/users/reactivate", json={
            "gmail": "lc_revoke@gmail.com"
        }, headers=_coord_headers())
        assert res.status_code == 200
        assert "already" in res.json()["message"].lower()

    def test_reactivate_invited_user_returns_error(self):
        db.grant_single_user_access("lc_invited_react@gmail.com", "Student")
        res = client.patch("/api/users/reactivate", json={
            "gmail": "lc_invited_react@gmail.com"
        }, headers=_coord_headers())
        assert res.status_code == 400
        assert "invited" in res.json()["message"].lower()

    def test_reactivate_nonexistent_404(self):
        res = client.patch("/api/users/reactivate", json={
            "gmail": "nonexistent_react@gmail.com"
        }, headers=_coord_headers())
        assert res.status_code == 404


# ==============================================================
# 35K: RESEND INVITATION
# ==============================================================

class TestResendInvitation:
    def test_resend_invitation_for_invited_user(self):
        db.grant_single_user_access("resend_test@gmail.com", "Student")
        res = client.post("/api/users/resend-invitation", json={
            "gmail": "resend_test@gmail.com"
        }, headers=_coord_headers())
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True
        assert "email_sent" in data

    def test_resend_invalidates_old_token(self):
        result = db.grant_single_user_access("resend_old@gmail.com", "Student")
        old_token = result["activation_token"]

        user, new_token = db.resend_invitation("resend_old@gmail.com")
        assert new_token is not None

        old_check = db.validate_auth_token(old_token, "ACTIVATION")
        assert not old_check["valid"]

        new_check = db.validate_auth_token(new_token, "ACTIVATION")
        assert new_check["valid"]

    def test_resend_creates_history(self):
        history = db.get_access_history(user_gmail="resend_test@gmail.com")
        resent = [h for h in history if h["action"] == "INVITATION_RESENT"]
        assert len(resent) >= 1

    def test_resend_for_active_user_fails(self):
        _create_active_user("resend_active@gmail.com", "Student")
        res = client.post("/api/users/resend-invitation", json={
            "gmail": "resend_active@gmail.com"
        }, headers=_coord_headers())
        assert res.status_code == 400

    def test_resend_for_nonexistent_user_fails(self):
        res = client.post("/api/users/resend-invitation", json={
            "gmail": "nonexistent_resend@gmail.com"
        }, headers=_coord_headers())
        assert res.status_code == 400

    def test_student_cannot_resend(self):
        student = db.get_user_by_gmail("student@gmail.com")
        res = client.post("/api/users/resend-invitation", json={
            "gmail": "resend_test@gmail.com"
        }, headers={"X-User-Id": student["uuid"]})
        assert res.status_code == 403


# ==============================================================
# 35L: BULK INVITATION
# ==============================================================

class TestBulkInvitation:
    def test_bulk_creates_invited_users(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["gmail", "role"])
        ws.append(["bulk_lc1@gmail.com", "Student"])
        ws.append(["bulk_lc2@gmail.com", "Mentor"])

        f = io.BytesIO()
        wb.save(f)
        f.seek(0)

        res = client.post("/api/users/upload-access",
            data={"default_role": "Student"},
            files={"file": ("test.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=_coord_headers()
        )
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True
        assert data["invited_count"] == 2

        u1 = db.get_user_by_gmail("bulk_lc1@gmail.com")
        assert u1["access_status"] == "INVITED"
        assert u1["is_active"] == 0

        u2 = db.get_user_by_gmail("bulk_lc2@gmail.com")
        assert u2["access_status"] == "INVITED"

    def test_bulk_skips_revoked_users(self):
        _create_active_user("bulk_skip_rev@gmail.com", "Student")
        db.revoke_user_access("bulk_skip_rev@gmail.com")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["gmail", "role"])
        ws.append(["bulk_skip_rev@gmail.com", "Student"])

        f = io.BytesIO()
        wb.save(f)
        f.seek(0)

        res = client.post("/api/users/upload-access",
            data={"default_role": "Student"},
            files={"file": ("test2.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=_coord_headers()
        )
        assert res.status_code == 200
        data = res.json()
        assert data["skipped_revoked"] >= 1

        user = db.get_user_by_gmail("bulk_skip_rev@gmail.com")
        assert user["is_active"] == 0

    def test_bulk_skips_already_invited(self):
        db.grant_single_user_access("bulk_skip_inv@gmail.com", "Student")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["gmail", "role"])
        ws.append(["bulk_skip_inv@gmail.com", "Student"])

        f = io.BytesIO()
        wb.save(f)
        f.seek(0)

        res = client.post("/api/users/upload-access",
            data={"default_role": "Student"},
            files={"file": ("test3.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=_coord_headers()
        )
        data = res.json()
        assert data["skipped_invited"] >= 1


# ==============================================================
# 35M: AUTHORIZATION
# ==============================================================

class TestAuthorization:
    def test_missing_auth_header_rejected(self):
        res = client.get("/api/users/managed")
        assert res.status_code == 401

    def test_unknown_user_rejected(self):
        res = client.get("/api/users/managed", headers={"X-User-Id": "fake-uuid-999"})
        assert res.status_code == 401

    def test_student_cannot_access_managed(self):
        student = db.get_user_by_gmail("student@gmail.com")
        res = client.get("/api/users/managed", headers={"X-User-Id": student["uuid"]})
        assert res.status_code == 403

    def test_revoked_coordinator_blocked(self):
        _create_active_user("auth_revoked_coord@gmail.com", "Coordinator")
        db.revoke_user_access("auth_revoked_coord@gmail.com")
        user = db.get_user_by_gmail("auth_revoked_coord@gmail.com")
        res = client.get("/api/users/managed", headers={"X-User-Id": user["uuid"]})
        assert res.status_code == 403


# ==============================================================
# 35N: FULL END-TO-END LIFECYCLE
# ==============================================================

class TestFullLifecycle:
    @patch.dict(os.environ, {"SMTP_HOST": "", "SMTP_USERNAME": "", "SMTP_PASSWORD": "", "SMTP_FROM": ""}, clear=False)
    def test_complete_access_lifecycle(self):
        """
        Full flow: invite → activation blocked login → activate → login →
        revoke → login blocked → reactivate → login → forgot password → reset → login
        """
        gmail = "e2e_lifecycle@gmail.com"

        # 1. Coordinator invites user
        res = client.post("/api/users/grant-single-access", json={
            "gmail": gmail, "role": "Student"
        }, headers=_coord_headers())
        assert res.status_code == 200
        data = res.json()
        assert data["success"] is True
        activation_url = data.get("dev_activation_url", "") or data["user"].get("dev_activation_url", "")
        assert "#/activate/" in activation_url, f"No activation URL found. Keys: {list(data.keys())}"
        token = activation_url.split("#/activate/")[1]

        # 2. INVITED user cannot login
        login_res = client.post("/api/login", json={"gmail": gmail, "password": "any"})
        assert login_res.status_code == 403

        # 3. User activates account
        act_res = client.post("/api/auth/activate", json={
            "token": token, "password": "E2ePass123"
        })
        assert act_res.status_code == 200

        user = db.get_user_by_gmail(gmail)
        assert user["access_status"] == "ACTIVE"
        assert user["is_active"] == 1
        assert db.is_hashed(user["password"])

        # 4. User can now login
        login_res = client.post("/api/login", json={"gmail": gmail, "password": "E2ePass123"})
        assert login_res.status_code == 200

        # 5. Coordinator revokes
        rev_res = client.patch("/api/users/revoke", json={"gmail": gmail}, headers=_coord_headers())
        assert rev_res.status_code == 200

        # 6. Login blocked
        login_res = client.post("/api/login", json={"gmail": gmail, "password": "E2ePass123"})
        assert login_res.status_code == 403

        # 7. Coordinator reactivates
        react_res = client.patch("/api/users/reactivate", json={"gmail": gmail}, headers=_coord_headers())
        assert react_res.status_code == 200

        # 8. Login works again
        login_res = client.post("/api/login", json={"gmail": gmail, "password": "E2ePass123"})
        assert login_res.status_code == 200

        # 9. Forgot password flow
        forgot_res = client.post("/api/auth/forgot-password", json={"gmail": gmail})
        assert forgot_res.status_code == 200
        reset_url = forgot_res.json().get("dev_reset_url", "")
        assert "#/reset-password/" in reset_url
        reset_token = reset_url.split("#/reset-password/")[1]

        # 10. Reset password
        reset_res = client.post("/api/auth/reset-password", json={
            "token": reset_token, "password": "NewE2ePass456"
        })
        assert reset_res.status_code == 200

        # 11. Login with new password
        login_res = client.post("/api/login", json={"gmail": gmail, "password": "NewE2ePass456"})
        assert login_res.status_code == 200

        # 12. Old password fails
        login_res = client.post("/api/login", json={"gmail": gmail, "password": "E2ePass123"})
        assert login_res.status_code == 401

        # 13. Verify full audit trail
        history = db.get_access_history(user_gmail=gmail)
        actions = [h["action"] for h in history]
        assert "INVITED" in actions
        assert "ACTIVATED" in actions
        assert "REVOKED" in actions
        assert "REACTIVATED" in actions
        assert "PASSWORD_RESET" in actions


# ==============================================================
# 35O: EMAIL DELIVERY — MOCKED SMTP
# ==============================================================

_SMTP_ENV = {
    "SMTP_HOST": "smtp.gmail.com",
    "SMTP_PORT": "587",
    "SMTP_USERNAME": "test@gmail.com",
    "SMTP_PASSWORD": "app-password",
    "SMTP_FROM": "test@gmail.com",
    "SMTP_USE_TLS": "true",
}

_SMTP_ENV_NO_TLS = {**_SMTP_ENV, "SMTP_USE_TLS": "false"}

_SMTP_ENV_EMPTY = {
    "SMTP_HOST": "",
    "SMTP_USERNAME": "",
    "SMTP_PASSWORD": "",
    "SMTP_FROM": "",
}

_SMTP_ENV_PARTIAL = {
    "SMTP_HOST": "smtp.gmail.com",
    "SMTP_USERNAME": "",
    "SMTP_PASSWORD": "",
    "SMTP_FROM": "test@gmail.com",
}


class TestEmailDeliverySmtpConfigured:
    """Tests with SMTP fully configured and delivery succeeding."""

    @patch("email_service.smtplib.SMTP")
    @patch.dict(os.environ, _SMTP_ENV)
    def test_invitation_email_sent_via_smtp(self, mock_smtp_cls):
        mock_server = MagicMock()
        mock_smtp_cls.return_value.__enter__ = MagicMock(return_value=mock_server)
        mock_smtp_cls.return_value.__exit__ = MagicMock(return_value=False)

        result = email_service.send_invitation_email(
            "recipient@gmail.com", "Student", "http://localhost:8000/#/activate/token123"
        )
        assert result["sent"] is True
        assert result["delivery_mode"] == "smtp"
        mock_server.sendmail.assert_called_once()
        args = mock_server.sendmail.call_args[0]
        assert args[0] == "test@gmail.com"
        assert args[1] == ["recipient@gmail.com"]

    @patch("email_service.smtplib.SMTP")
    @patch.dict(os.environ, _SMTP_ENV)
    def test_password_reset_email_sent_via_smtp(self, mock_smtp_cls):
        mock_server = MagicMock()
        mock_smtp_cls.return_value.__enter__ = MagicMock(return_value=mock_server)
        mock_smtp_cls.return_value.__exit__ = MagicMock(return_value=False)

        result = email_service.send_password_reset_email(
            "user@gmail.com", "http://localhost:8000/#/reset-password/token456"
        )
        assert result["sent"] is True
        assert result["delivery_mode"] == "smtp"
        mock_server.starttls.assert_called_once()
        mock_server.login.assert_called_once_with("test@gmail.com", "app-password")

    @patch("email_service.smtplib.SMTP")
    @patch.dict(os.environ, _SMTP_ENV_NO_TLS)
    def test_smtp_without_tls(self, mock_smtp_cls):
        mock_server = MagicMock()
        mock_smtp_cls.return_value.__enter__ = MagicMock(return_value=mock_server)
        mock_smtp_cls.return_value.__exit__ = MagicMock(return_value=False)

        result = email_service.send_invitation_email(
            "notls@gmail.com", "Student", "http://localhost:8000/#/activate/tok"
        )
        assert result["sent"] is True
        mock_server.starttls.assert_not_called()

    @patch.dict(os.environ, {**_SMTP_ENV, "SMTP_PASSWORD": "xxxx xxxx xxxx xxxx"})
    def test_app_password_spaces_normalized(self):
        cfg = email_service._get_smtp_config()
        assert " " not in cfg["password"]
        assert cfg["password"] == "xxxxxxxxxxxxxxxx"


class TestEmailDeliverySmtpFailure:
    """Tests where SMTP is configured but delivery fails."""

    @patch("email_service.smtplib.SMTP")
    @patch.dict(os.environ, {**_SMTP_ENV, "SMTP_PASSWORD": "wrong-password"})
    def test_smtp_auth_failure(self, mock_smtp_cls):
        mock_server = MagicMock()
        mock_server.login.side_effect = smtplib.SMTPAuthenticationError(535, b"Auth failed")
        mock_smtp_cls.return_value.__enter__ = MagicMock(return_value=mock_server)
        mock_smtp_cls.return_value.__exit__ = MagicMock(return_value=False)

        result = email_service.send_invitation_email(
            "user@gmail.com", "Student", "http://localhost:8000/#/activate/tok"
        )
        assert result["sent"] is False
        assert result["delivery_mode"] == "smtp"
        assert "authentication" in result["reason"].lower()

    @patch("email_service.smtplib.SMTP")
    @patch.dict(os.environ, _SMTP_ENV)
    def test_smtp_generic_error(self, mock_smtp_cls):
        mock_server = MagicMock()
        mock_server.sendmail.side_effect = smtplib.SMTPException("Connection lost")
        mock_smtp_cls.return_value.__enter__ = MagicMock(return_value=mock_server)
        mock_smtp_cls.return_value.__exit__ = MagicMock(return_value=False)

        result = email_service.send_password_reset_email(
            "user@gmail.com", "http://localhost:8000/#/reset-password/tok"
        )
        assert result["sent"] is False
        assert result["delivery_mode"] == "smtp"
        assert "smtp error" in result["reason"].lower()

    @patch("email_service.smtplib.SMTP")
    @patch.dict(os.environ, _SMTP_ENV)
    def test_smtp_unexpected_exception(self, mock_smtp_cls):
        mock_server = MagicMock()
        mock_server.sendmail.side_effect = OSError("Network unreachable")
        mock_smtp_cls.return_value.__enter__ = MagicMock(return_value=mock_server)
        mock_smtp_cls.return_value.__exit__ = MagicMock(return_value=False)

        result = email_service.send_invitation_email(
            "user@gmail.com", "Student", "http://localhost:8000/#/activate/tok"
        )
        assert result["sent"] is False
        assert "failed" in result["reason"].lower()


class TestEmailDeliveryDevFallback:
    """Tests where SMTP is NOT configured — development fallback."""

    @patch.dict(os.environ, _SMTP_ENV_EMPTY, clear=False)
    def test_dev_fallback_invitation(self):
        result = email_service.send_invitation_email(
            "dev@gmail.com", "Student", "http://localhost:8000/#/activate/tok"
        )
        assert result["sent"] is False
        assert result["delivery_mode"] == "development"
        assert "not configured" in result["reason"].lower()

    @patch.dict(os.environ, _SMTP_ENV_EMPTY, clear=False)
    def test_dev_fallback_password_reset(self):
        result = email_service.send_password_reset_email(
            "dev@gmail.com", "http://localhost:8000/#/reset-password/tok"
        )
        assert result["sent"] is False
        assert result["delivery_mode"] == "development"

    @patch.dict(os.environ, _SMTP_ENV_PARTIAL, clear=False)
    def test_partial_smtp_config_falls_back(self):
        result = email_service.send_invitation_email(
            "partial@gmail.com", "Student", "http://localhost:8000/#/activate/tok"
        )
        assert result["sent"] is False
        assert result["delivery_mode"] == "development"


class TestTruthfulApiEmailResponse:
    """Tests that API endpoints return truthful email_sent / delivery_mode fields."""

    @patch("email_service.smtplib.SMTP")
    @patch.dict(os.environ, _SMTP_ENV)
    def test_grant_access_email_sent_true(self, mock_smtp_cls):
        mock_server = MagicMock()
        mock_smtp_cls.return_value.__enter__ = MagicMock(return_value=mock_server)
        mock_smtp_cls.return_value.__exit__ = MagicMock(return_value=False)

        res = client.post("/api/users/grant-single-access", json={
            "gmail": "email_test_grant@gmail.com", "role": "Student"
        }, headers=_coord_headers())
        data = res.json()
        assert data["email_sent"] is True
        assert data["delivery_mode"] == "smtp"
        assert "sent" in data["message"].lower()
        assert "dev_activation_url" not in data

    @patch.dict(os.environ, _SMTP_ENV_EMPTY, clear=False)
    def test_grant_access_dev_fallback(self):
        res = client.post("/api/users/grant-single-access", json={
            "gmail": "email_test_dev@gmail.com", "role": "Student"
        }, headers=_coord_headers())
        data = res.json()
        assert data["email_sent"] is False
        assert data["delivery_mode"] == "development"
        assert "not configured" in data["message"].lower()
        assert "dev_activation_url" in data

    @patch("email_service.smtplib.SMTP")
    @patch.dict(os.environ, {**_SMTP_ENV, "SMTP_PASSWORD": "bad"})
    def test_grant_access_smtp_failure(self, mock_smtp_cls):
        mock_server = MagicMock()
        mock_server.login.side_effect = smtplib.SMTPAuthenticationError(535, b"Bad creds")
        mock_smtp_cls.return_value.__enter__ = MagicMock(return_value=mock_server)
        mock_smtp_cls.return_value.__exit__ = MagicMock(return_value=False)

        res = client.post("/api/users/grant-single-access", json={
            "gmail": "email_test_fail@gmail.com", "role": "Student"
        }, headers=_coord_headers())
        data = res.json()
        assert data["email_sent"] is False
        assert "could not be delivered" in data["message"].lower()
        assert "dev_activation_url" in data

    @patch("email_service.smtplib.SMTP")
    @patch.dict(os.environ, _SMTP_ENV)
    def test_resend_invitation_email_sent(self, mock_smtp_cls):
        mock_server = MagicMock()
        mock_smtp_cls.return_value.__enter__ = MagicMock(return_value=mock_server)
        mock_smtp_cls.return_value.__exit__ = MagicMock(return_value=False)

        db.grant_single_user_access("email_resend_test@gmail.com", "Student")
        res = client.post("/api/users/resend-invitation", json={
            "gmail": "email_resend_test@gmail.com"
        }, headers=_coord_headers())
        data = res.json()
        assert data["email_sent"] is True
        assert data["delivery_mode"] == "smtp"
        assert "resent" in data["message"].lower()

    @patch.dict(os.environ, _SMTP_ENV_EMPTY, clear=False)
    def test_resend_invitation_dev_fallback(self):
        db.grant_single_user_access("email_resend_dev@gmail.com", "Student")
        res = client.post("/api/users/resend-invitation", json={
            "gmail": "email_resend_dev@gmail.com"
        }, headers=_coord_headers())
        data = res.json()
        assert data["email_sent"] is False
        assert data["delivery_mode"] == "development"
        assert "not configured" in data["message"].lower()
        assert "dev_activation_url" in data

    @patch.dict(os.environ, _SMTP_ENV_EMPTY, clear=False)
    def test_forgot_password_dev_fallback(self):
        _create_active_user("email_forgot_test@gmail.com", "Student")
        res = client.post("/api/auth/forgot-password", json={
            "gmail": "email_forgot_test@gmail.com"
        })
        data = res.json()
        assert data["email_sent"] is False
        assert data["delivery_mode"] == "development"
        assert "dev_reset_url" in data

    @patch("email_service.smtplib.SMTP")
    @patch.dict(os.environ, _SMTP_ENV)
    def test_forgot_password_email_sent(self, mock_smtp_cls):
        mock_server = MagicMock()
        mock_smtp_cls.return_value.__enter__ = MagicMock(return_value=mock_server)
        mock_smtp_cls.return_value.__exit__ = MagicMock(return_value=False)

        _create_active_user("email_forgot_smtp@gmail.com", "Student")
        res = client.post("/api/auth/forgot-password", json={
            "gmail": "email_forgot_smtp@gmail.com"
        })
        data = res.json()
        assert data["email_sent"] is True
        assert data["delivery_mode"] == "smtp"
        assert "dev_reset_url" not in data


# ==============================================================
# REGRESSION: EXISTING FUNCTIONALITY
# ==============================================================

class TestRegression:
    def test_seed_coordinator_login(self):
        res = client.post("/api/login", json={
            "gmail": "coordinator@gmail.com", "password": "coord123"
        })
        assert res.status_code == 200
        assert res.json()["user"]["role"] == "Coordinator"

    def test_seed_student_login(self):
        res = client.post("/api/login", json={
            "gmail": "student@gmail.com", "password": "student123"
        })
        assert res.status_code == 200

    def test_drives_endpoint(self):
        res = client.get("/api/drives")
        assert res.status_code == 200
        assert res.json()["success"] is True

    def test_students_endpoint(self):
        res = client.get("/api/students")
        assert res.status_code == 200

    def test_users_endpoint(self):
        res = client.get("/api/users")
        assert res.status_code == 200

    def test_invalid_login_rejected(self):
        res = client.post("/api/login", json={
            "gmail": "coordinator@gmail.com", "password": "wrong"
        })
        assert res.status_code == 401

    def test_nonexistent_user_login(self):
        res = client.post("/api/login", json={
            "gmail": "nobody@gmail.com", "password": "pass"
        })
        assert res.status_code == 401
